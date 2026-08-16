"""Annexure-IV Table-1 / Table-2 exports in the exact prescribed layout
(SB Order 09/2026, pages 13-14)."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(name="Arial", bold=True)
BODY = Font(name="Arial")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
GREY = PatternFill("solid", fgColor="EFEFEF")
MONEY_FMT = "#,##0.00"


def _title_block(ws, title_lines: list[str], ncols: int) -> int:
    row = 1
    for line in title_lines:
        cell = ws.cell(row=row, column=1, value=line)
        cell.font = Font(name="Arial", bold=True, size=12)
        cell.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        row += 1
    return row + 1  # blank spacer row


def _ddo_month_row(ws, row: int, ddo_code: str, month_label: str, ncols: int) -> int:
    left = ws.cell(row=row, column=1, value=f"DDO Code: {ddo_code}")
    left.font = BOLD
    right = ws.cell(row=row, column=ncols, value=f"Month: {month_label}")
    right.font = BOLD
    right.alignment = Alignment(horizontal="right")
    return row + 1


def _two_row_header(ws, row: int, groups: list[tuple[str, list[str]]]) -> int:
    """Write the merged two-row header. groups = [(top_label, [sub_labels...])]."""
    col = 1
    for top, subs in groups:
        span = max(len(subs), 1)
        cell = ws.cell(row=row, column=col, value=top)
        cell.font = BOLD
        cell.alignment = CENTER
        cell.fill = GREY
        if not subs:  # single column spanning both header rows
            ws.merge_cells(start_row=row, start_column=col, end_row=row + 1, end_column=col)
            ws.cell(row=row + 1, column=col).border = BORDER
        else:
            if span > 1:
                ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
            for i, sub in enumerate(subs):
                sub_cell = ws.cell(row=row + 1, column=col + i, value=sub)
                sub_cell.font = BOLD
                sub_cell.alignment = CENTER
                sub_cell.fill = GREY
                sub_cell.border = BORDER
        for i in range(span):
            ws.cell(row=row, column=col + i).border = BORDER
        col += span
    return row + 2


def _data_rows(ws, row: int, df: pd.DataFrame, money_from: int) -> int:
    for record in df.itertuples(index=False):
        for c, value in enumerate(record, start=1):
            cell = ws.cell(row=row, column=c, value=value)
            cell.font = BODY
            cell.border = BORDER
            if c >= money_from:
                cell.number_format = MONEY_FMT
        row += 1
    return row


def _total_row(ws, row: int, df: pd.DataFrame, money_from: int, ncols: int) -> int:
    cell = ws.cell(row=row, column=1, value="TOTAL")
    cell.font = BOLD
    cell.fill = GREY
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=money_from - 1)
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).border = BORDER
        ws.cell(row=row, column=c).fill = GREY
    for c in range(money_from, ncols + 1):
        col_name = df.columns[c - 1]
        total = float(df[col_name].fillna(0).sum()) if len(df) else 0.0
        cell = ws.cell(row=row, column=c, value=round(total, 2))
        cell.font = BOLD
        cell.number_format = MONEY_FMT
    return row + 1


def _signature_block(ws, row: int, ho_name: str, division: str, ncols: int) -> None:
    row += 2
    left = ws.cell(row=row, column=1, value="SBCO In-charge")
    left.font = BOLD
    right = ws.cell(row=row, column=ncols, value="Postmaster")
    right.font = BOLD
    right.alignment = Alignment(horizontal="right")
    row += 1
    ws.cell(row=row, column=1, value=f"{ho_name or '____________'} HO").font = BODY
    r2 = ws.cell(row=row, column=ncols, value=f"{ho_name or '____________'} HO")
    r2.font = BODY
    r2.alignment = Alignment(horizontal="right")
    row += 2
    ws.cell(row=row, column=1, value="Forwarded to the General Manager(F) / DA(P)").font = BODY
    row += 1
    ws.cell(
        row=row, column=1,
        value=f"Copy to: The S/SPOs, {division or '________________'} Division for information.",
    ).font = BODY


def _autosize(ws, ncols: int) -> None:
    widths = {1: 5, 2: 14, 3: 40}
    for c in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 14)


def month_label(month: str) -> str:
    """'2026-07' -> 'July 2026'."""
    date = dt.date(int(month[:4]), int(month[5:7]), 1)
    return date.strftime("%B %Y")


def export_table1(
    df: pd.DataFrame, path: str | Path, month: str, ddo_code: str, ho_name: str, division: str,
) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Table-1"
    ncols = 9

    row = _title_block(ws, ["CBS Monthly Reconciliation Report to PAO by the HO",
                            "(Due Date: 4th of every month)", "Table-1"], ncols)
    row = _ddo_month_row(ws, row, ddo_code, month_label(month), ncols)
    row = _two_row_header(
        ws, row,
        [
            ("Sl", []),
            ("Account Code", []),
            ("Account Code Description", []),
            ("Finacle", ["Receipts\n(In Rs.)", "Payments\n(In Rs.)"]),
            ("Monthly Cash Account*", ["Receipts\n(In Rs.)", "Payments\n(In Rs.)"]),
            ("Difference\n(Finacle – Cash Account)", ["Receipts\n(In Rs.)", "Payments\n(In Rs.)"]),
        ],
    )
    row = _data_rows(ws, row, df, money_from=4)
    row = _total_row(ws, row, df, money_from=4, ncols=ncols)
    row += 1
    note = ws.cell(
        row=row, column=1,
        value="* Monthly Cash Account = Sum of Daily Cash Books + Approved Transfer Entries of DDO",
    )
    note.font = Font(name="Arial", italic=True)
    _signature_block(ws, row, ho_name, division, ncols)
    _autosize(ws, ncols)
    ws.freeze_panes = "A8"
    wb.save(path)
    return path


def export_table2(
    df: pd.DataFrame, path: str | Path, month: str, ddo_code: str, ho_name: str, division: str,
    pending_reasons: str = "", pendency_clear_date: str = "",
) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Table-2"
    ncols = 11

    row = _title_block(ws, ["Detailed CBS Monthly Reconciliation Report to PAO by the HO",
                            "(Due Date: 4th of every month)", "Table-2"], ncols)
    row = _ddo_month_row(ws, row, ddo_code, month_label(month), ncols)
    rp = ["Receipts\n(In Rs.)", "Payments\n(In Rs.)"]
    row = _two_row_header(
        ws, row,
        [
            ("Sl", []),
            ("Account Code", []),
            ("Account Code Description", []),
            ("Opening Balance in the Difference\nFinacle – Cash Account", rp),
            ("Current Month Difference\n(Finacle – Cash Account)", rp),
            ("Rectified During the Current Month\n(Finacle – Cash Account)", rp),
            ("Pending for Rectification", rp),
        ],
    )
    row = _data_rows(ws, row, df, money_from=4)
    row = _total_row(ws, row, df, money_from=4, ncols=ncols)
    row += 1
    ws.cell(row=row, column=1, value=f"(a) Reasons for pending for rectification : {pending_reasons}").font = BODY
    row += 1
    ws.cell(row=row, column=1, value=f"(b) Date by which the pendency is cleared : {pendency_clear_date}").font = BODY
    _signature_block(ws, row, ho_name, division, ncols)
    _autosize(ws, ncols)
    ws.freeze_panes = "A8"
    wb.save(path)
    return path
