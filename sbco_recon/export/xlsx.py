"""XLSX export of reconciliation views and the Discrepancy Register."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(name="Arial", bold=True)
BODY_FONT = Font(name="Arial")


def export_frame(
    df: pd.DataFrame,
    path: str | Path,
    title: str,
    subtitle: str = "",
    money_columns: list[str] | None = None,
) -> Path:
    """Write a DataFrame as a formatted report sheet."""
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    ncols = max(len(df.columns), 1)
    ws.cell(row=1, column=1, value=title).font = Font(name="Arial", bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    row_offset = 3
    if subtitle:
        ws.cell(row=2, column=1, value=subtitle).font = Font(name="Arial", italic=True)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        row_offset = 4

    money = set(money_columns or [])
    for c, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=row_offset, column=c, value=str(col))
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for r, record in enumerate(df.itertuples(index=False), start=row_offset + 1):
        for c, (col, value) in enumerate(zip(df.columns, record), start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            if col in money:
                cell.number_format = "#,##0.00"

    for c, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(width, 50)

    ws.freeze_panes = ws.cell(row=row_offset + 1, column=1)
    wb.save(path)
    return path
