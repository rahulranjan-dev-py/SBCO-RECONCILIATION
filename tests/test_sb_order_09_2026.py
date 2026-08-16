"""Conformance with SB Order No. 09/2026, Annexure-IV.

The order prescribes three forms. These tests pin the column layouts and the
arithmetic to the order itself, so a later change cannot quietly drift the
returns away from the prescribed format.
"""

from __future__ import annotations

import sys
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from sbco_recon.annexure import (RegisterEntry, Table1Row, build_table1,
                                 entries_from_reconciliation,
                                 expand_transfer_entries)
from sbco_recon.reports import (write_annexure_iv_table1,
                                write_annexure_iv_table2,
                                write_discrepancy_register)
from sbco_recon.table2 import Table2Row, build_table2

D = Decimal
RECEIPT_CODE, PAYMENT_CODE = "8001000100", "8001000200"


@dataclass
class Recon:
    rows: list
    warnings: list = None

    def __post_init__(self):
        self.warnings = self.warnings or []


@dataclass
class ReconRow:
    account_code: str
    description: str
    finacle: Decimal
    cashbook: Decimal

    @property
    def difference(self):
        return self.finacle - self.cashbook

    @property
    def matched(self):
        return self.difference == 0


# ── the no-netting rule, para 1(ix) ──────────────────────────────────

def test_receipt_and_payment_discrepancies_are_never_netted():
    """The order's own example: 5,000 on receipts and 2,000 on payments are
    reported separately, never as a net 3,000."""
    result = build_table1(Recon([
        ReconRow(RECEIPT_CODE, "POSB Rec", D("5000"), ZERO := D("0")),
        ReconRow(PAYMENT_CODE, "POSB Pay", D("2000"), D("0")),
    ]), "Jul-2026")
    totals = result.totals()
    assert totals["difference_receipt"] == D("5000")
    assert totals["difference_payment"] == D("2000")
    by_code = {r.account_code: r for r in result.rows}
    assert by_code[RECEIPT_CODE].difference_payment == 0
    assert by_code[PAYMENT_CODE].difference_receipt == 0


# ── Table-1 ──────────────────────────────────────────────────────────

def test_table1_columns_match_the_order(tmp_path):
    result = build_table1(Recon([
        ReconRow(RECEIPT_CODE, "POSB Rec", D("1500"), D("1000"))]), "Jul-2026")
    out = write_annexure_iv_table1(tmp_path / "t1.xlsx", result, ddo_code="102617",
                                   ho_name="Manipal HO", division="Manipal",
                                   month=date(2026, 7, 1))
    ws = load_workbook(out).active

    assert ws["A1"].value == "CBS Monthly Reconciliation Report to PAO by the HO"
    assert ws["A2"].value == "(Due Date: 4th of every month)"
    assert ws["A4"].value == "Table-1"
    assert ws["A7"].value == "Sl"
    assert ws["B7"].value == "Account Code"
    assert ws["C7"].value == "Account Code Description"
    assert ws["D7"].value == "Finacle"
    assert ws["F7"].value == "Monthly Cash Account*"
    assert ws["H7"].value.startswith("Difference")
    for ref in ("D8", "F8", "H8"):
        assert ws[ref].value.startswith("Receipts")
    for ref in ("E8", "G8", "I8"):
        assert ws[ref].value.startswith("Payments")


def test_table1_carries_the_cash_account_footnote_and_joint_signatures(tmp_path):
    result = build_table1(Recon([
        ReconRow(PAYMENT_CODE, "POSB Pay", D("24"), D("0"))]), "Jul-2026")
    out = write_annexure_iv_table1(tmp_path / "t1.xlsx", result, ddo_code="1",
                                   ho_name="HO", division="Udupi",
                                   month=date(2026, 7, 1))
    text = " ".join(str(c.value) for row in load_workbook(out).active.iter_rows()
                    for c in row if c.value is not None)
    assert "Sum of Daily Cash Books + Approved Transfer Entries of DDO" in text
    assert "SBCO In-charge" in text, "para 4(ix) requires joint signature"
    assert "Postmaster" in text
    assert "Forwarded to the General Manager(F) / DA(P)" in text
    assert "Copy to: The S/SPOs, Udupi Division for information." in text


def test_approved_transfer_entries_go_into_the_monthly_cash_account():
    """The order defines Monthly Cash Account as daily cash books plus
    approved TEs of the DDO."""
    recon = Recon([ReconRow(PAYMENT_CODE, "POSB Pay", D("1269"), D("1000"))])
    plain = build_table1(recon, "Jul-2026")
    assert plain.rows[0].difference_payment == D("269")

    adjusted = build_table1(recon, "Jul-2026",
                            [{"from_code": RECEIPT_CODE, "to_code": PAYMENT_CODE,
                              "amount": "269"}])
    by_code = {r.account_code: r for r in adjusted.rows}
    assert by_code[PAYMENT_CODE].cashbook_payment == D("1269")
    assert by_code[PAYMENT_CODE].difference_payment == 0
    assert by_code[RECEIPT_CODE].cashbook_receipt == D("-269"), "other leg missing"


def test_a_transfer_entry_posts_both_legs():
    expanded = expand_transfer_entries(
        [{"from_code": "8001000100", "to_code": "8001000200", "amount": "269"}])
    assert sorted(expanded) == [("8001000100", D("-269")),
                                ("8001000200", D("269"))]


# ── Table-2 ──────────────────────────────────────────────────────────

def test_table2_columns_match_the_order(tmp_path):
    seed = Table2Row(PAYMENT_CODE, "POSB Pay")
    seed.current_payment = D("269")
    result = build_table2("Aug-2026", [], previous_table2=[seed])
    out = write_annexure_iv_table2(tmp_path / "t2.xlsx", result, ddo_code="1",
                                   ho_name="HO", division="Udupi",
                                   month=date(2026, 8, 1))
    ws = load_workbook(out).active

    assert ws["A1"].value == ("Detailed CBS Monthly Reconciliation Report "
                              "to PAO by the HO")
    assert ws["A4"].value == "Table-2"
    assert ws["D7"].value.startswith("Opening Balance")
    assert ws["F7"].value.startswith("Current Month")
    assert ws["H7"].value.startswith("Rectified During the Current Month")
    assert ws["J7"].value == "Pending for Rectification"


def test_table2_carries_the_two_required_footer_questions(tmp_path):
    result = build_table2("Aug-2026", [])
    out = write_annexure_iv_table2(tmp_path / "t2.xlsx", result, ddo_code="1",
                                   ho_name="HO", division="D",
                                   month=date(2026, 8, 1))
    text = " ".join(str(c.value) for row in load_workbook(out).active.iter_rows()
                    for c in row if c.value is not None)
    assert "(a) Reasons for pending for rectification" in text
    assert "(b) Date by which the pendency is cleared" in text
    assert "SBCO In-charge" in text and "Postmaster" in text


# ── Table-3, the daily register ──────────────────────────────────────

def test_table3_columns_match_the_order(tmp_path):
    entry = RegisterEntry(
        entry_date=date(2026, 7, 18), account_code=PAYMENT_CODE,
        office_name="Manipal SO", cbs_payment=D("1066791"),
        cashbook_payment=D("1493816"), sbco_initials="SK")
    out = write_discrepancy_register(tmp_path / "t3.xlsx", [entry],
                                     financial_year="2026/27", ho_name="Manipal HO")
    ws = load_workbook(out).active
    headers = [ws.cell(row=5, column=c).value for c in range(1, 18)]
    assert headers[0] == "Sl. No."
    assert headers[2] == "Account Code"
    assert headers[4].startswith("Name of the Post office")
    assert headers[9].startswith("Difference (CBS - Cashbook)")
    assert headers[11] == "Initials of In-Charge SBCO"
    assert headers[12] == "Date of Rectification"
    assert headers[13].startswith("Particulars of Misc. Transaction")
    assert headers[14].startswith("Particulars of Transfer Entries")
    assert headers[16] == "Initials of Postmaster"
    # the order labels its columns (a)-(o); the difference columns are formulas
    assert ws.cell(row=6, column=10).value == "(f) - (h)"
    assert ws.cell(row=7, column=10).value == "=F7-H7"
    assert ws.cell(row=7, column=11).value == "=G7-I7"


def test_table3_states_the_financial_year_reset_rule(tmp_path):
    out = write_discrepancy_register(tmp_path / "t3.xlsx", [],
                                     financial_year="2026/27")
    text = " ".join(str(c.value) for row in load_workbook(out).active.iter_rows()
                    for c in row if c.value is not None)
    assert "closed at the end of each financial year" in text
    assert "reset from S No.1" in text


def test_register_entries_are_seeded_from_a_reconciliation():
    recon = Recon([
        ReconRow(RECEIPT_CODE, "POSB Rec", D("1500"), D("1000")),
        ReconRow(PAYMENT_CODE, "POSB Pay", D("100"), D("100")),      # matched
    ])
    entries = entries_from_reconciliation(recon, date(2026, 7, 18), "Manipal SO")
    assert len(entries) == 1, "matched codes must not enter the register"
    entry = entries[0]
    assert entry.account_code == RECEIPT_CODE
    assert entry.cbs_receipt == D("1500")
    assert entry.cbs_payment == 0
    assert entry.difference_receipt == D("500")
    assert entry.financial_year == "2026/27"
    assert not entry.is_settled


def test_a_register_entry_is_settled_only_once_rectification_is_recorded():
    entry = RegisterEntry(entry_date=date(2026, 7, 18), account_code=PAYMENT_CODE)
    assert not entry.is_settled
    entry.rectified_date = date(2026, 7, 21)
    assert entry.is_settled
    assert entry.days_outstanding == 3


@pytest.mark.parametrize("d,expected", [
    (date(2026, 4, 1), "2026/27"), (date(2027, 3, 31), "2026/27"),
    (date(2026, 3, 31), "2025/26"),
])
def test_register_financial_year_boundaries(d, expected):
    assert RegisterEntry(entry_date=d, account_code="8001000100").financial_year \
        == expected
