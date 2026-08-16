from __future__ import annotations

import datetime as dt

from openpyxl import load_workbook

from sbco_recon.export import export_table1, export_table2
from sbco_recon.services import import_service, monthly_service, register_service

from .fixtures import make_cashbook_file, make_glwise_file

JULY = "2026-07"
AUG = "2026-08"
D_JUL = dt.date(2026, 7, 31)
D_AUG = dt.date(2026, 8, 14)


def _load_july(session, tmp_path):
    # Finacle: receipts 10000 (R-side code), payments 4000 (P-side code)
    make_glwise_file(
        tmp_path / "gl_jul.xlsx", D_JUL,
        [("8001000100", "POSB -Receipts", 10000.0, 0), ("8001000200", "POSB -Payments", 0, 4000.0)],
    )
    # Cashbook: receipts 9400 (600 short), payments 4000 (tallies)
    make_cashbook_file(
        tmp_path / "cb_jul.xlsx", D_JUL,
        [
            ("8001000100", "POSB -Receipts", "Receipts", 9400.0, 0, 0),
            ("8001000200", "POSB -Payments", "Payments", 4000.0, 0, 0),
        ],
    )
    results = import_service.import_files(session, [tmp_path / "gl_jul.xlsx", tmp_path / "cb_jul.xlsx"])
    assert all(r.status == "PROCESSED" for r in results)


def test_table1_rp_split_and_te(session, tmp_path):
    _load_july(session, tmp_path)
    # Approved TE of 500 TO the receipts code: Monthly Cash Account = cashbook + TE
    monthly_service.add_transfer_entry(session, JULY, "8001000100", 500.0, direction=+1)

    df = monthly_service.table1(session, JULY)
    r100 = df[df["Account Code"] == "8001000100"].iloc[0]
    # receipt-side code lands in the Receipts columns
    assert r100["Finacle Receipts"] == 10000.0 and r100["Finacle Payments"] == 0.0
    assert r100["Cash Account Receipts"] == 9900.0  # 9400 + 500 TE
    assert r100["Difference Receipts"] == 100.0

    r200 = df[df["Account Code"] == "8001000200"].iloc[0]
    # payment-side code lands in the Payments columns and tallies
    assert r200["Finacle Payments"] == 4000.0 and r200["Finacle Receipts"] == 0.0
    assert r200["Difference Payments"] == 0.0


def test_table2_carry_forward(session, tmp_path):
    _load_july(session, tmp_path)

    # August data: same codes, receipts differ by 200 this month
    make_glwise_file(tmp_path / "gl_aug.xlsx", D_AUG, [("8001000100", "POSB -Receipts", 5000.0, 0)])
    make_cashbook_file(
        tmp_path / "cb_aug.xlsx", D_AUG, [("8001000100", "POSB -Receipts", "Receipts", 4800.0, 0, 0)]
    )
    assert all(
        r.status == "PROCESSED"
        for r in import_service.import_files(session, [tmp_path / "gl_aug.xlsx", tmp_path / "cb_aug.xlsx"])
    )

    # July discrepancy (600 receipts) recorded and rectified during August
    entry = register_service.add_entry(
        session, D_JUL, "8001000100", cbs_receipt=10000.0, cb_receipt=9400.0
    )
    register_service.settle_entry(session, entry.id, rectified_on=dt.date(2026, 8, 5))

    df = monthly_service.table2(session, AUG)
    row = df[df["Account Code"] == "8001000100"].iloc[0]
    assert row["Opening Receipts"] == 600.0     # July difference carried forward
    assert row["Current Month Receipts"] == 200.0
    assert row["Rectified Receipts"] == 600.0   # settled in August
    assert row["Pending Receipts"] == 200.0     # 600 + 200 - 600


def test_annexure_exports(session, tmp_path):
    _load_july(session, tmp_path)
    t1 = monthly_service.table1(session, JULY)
    t2 = monthly_service.table2(session, JULY)

    p1 = export_table1(t1, tmp_path / "t1.xlsx", JULY, "102617", "Manipal", "Udupi")
    p2 = export_table2(t2, tmp_path / "t2.xlsx", JULY, "102617", "Manipal", "Udupi",
                       pending_reasons="Awaiting SO error book", pendency_clear_date="04/09/2026")

    ws = load_workbook(p1).active
    text = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("CBS Monthly Reconciliation Report" in t for t in text)
    assert any("DDO Code: 102617" in t for t in text)
    assert any("Month: July 2026" in t for t in text)
    assert any("TOTAL" == t for t in text)
    assert any("Monthly Cash Account = Sum of Daily Cash Books" in t for t in text)

    ws2 = load_workbook(p2).active
    text2 = [str(c.value) for row in ws2.iter_rows() for c in row if c.value is not None]
    assert any("Detailed CBS Monthly Reconciliation Report" in t for t in text2)
    assert any("Pending for Rectification" in t for t in text2)
    assert any("Awaiting SO error book" in t for t in text2)


def test_transfer_entry_crud(session):
    entry = monthly_service.add_transfer_entry(session, JULY, "8001000100", 250.0, direction=-1, remarks="reversal")
    df = monthly_service.transfer_entries_frame(session, JULY)
    assert len(df) == 1
    assert df.iloc[0]["To(+)/From(−)"] == "FROM (−)"
    assert monthly_service.delete_transfer_entry(session, entry.id)
    assert monthly_service.transfer_entries_frame(session, JULY).empty
