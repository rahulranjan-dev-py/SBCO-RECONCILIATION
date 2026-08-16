"""Builders for synthetic Finacle / APT report files.

The layouts replicate what the legacy VBA expected (cell landmarks, column order)
and the report sample shown in SB Order 09/2026 Annexure-III. Real anonymized
exports should be added to tests/fixtures/ as they become available and get their
own tests — these builders exist so the parsers are regression-tested from day one.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook

GLWISE_TITLE = "GL IT2.0 Transaction GL Wise Report (Incl HO,SO &BOs) - Consolidated(Previous Day)"


def make_glwise_file(
    path: Path,
    date: dt.date,
    rows: list[tuple[str, str, float, float]],  # (code, desc, deposits, withdrawals)
) -> Path:
    """Build a synthetic GL-wise consolidated report.

    Mirrors the real export shape: banner rows, labelled header block (Date in the
    F-ish column area, 'Set ID / Desc :' block), then the data table with
    S.No | GL Sub Head Code | IT2.0 A/C Code | IT2.0 Acct Code Desc | Deposits | Withdrawals.
    """
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "India Post"
    ws["I4"] = GLWISE_TITLE
    ws["N4"] = f"Run Date : {dt.date.today():%d/%m/%Y} 11.25 AM"
    ws["E8"] = "Date :"
    ws["F8"] = date.strftime("%d-%m-%Y")
    ws["E9"] = "Sol ID :"
    ws["F9"] = "60001700"
    ws["B14"] = "Set ID / Desc :"
    ws["C14"] = "TAMI1 - TAMI1"
    ws["B15"] = "Sol ID / Desc :"
    ws["C15"] = "60001700 - Model H.O"

    header_row = 17
    headers = ["S.No", "GL Sub Head Code", "IT2.0 A/C Code", "IT2.0 Acct Code Desc", "Deposits (Cr)", "Withdrawals (Dr)"]
    for idx, header in enumerate(headers):
        ws.cell(row=header_row, column=9 + idx, value=header)  # table starts at column I

    r = header_row + 1
    for sno, (code, desc, dep, wd) in enumerate(rows, start=1):
        ws.cell(row=r, column=9, value=sno)
        ws.cell(row=r, column=10, value="30001")
        ws.cell(row=r, column=11, value=int(code))
        ws.cell(row=r, column=12, value=desc)
        ws.cell(row=r, column=13, value=dep)
        ws.cell(row=r, column=14, value=wd)
        r += 1
    ws.cell(row=r, column=12, value="Total")
    ws.cell(row=r, column=13, value=sum(x[2] for x in rows))
    ws.cell(row=r, column=14, value=sum(x[3] for x in rows))

    wb.save(path)
    return path


def make_glwise_setid_file(
    path: Path,
    date: dt.date,
    sections: list[tuple[str, list[tuple[str, str, float, float]]]],  # (sol_id, rows)
) -> Path:
    """Build a synthetic Set-ID GL-wise report: one 'Sol ID / Desc :' block and
    data table per SOL, all in one sheet."""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "India Post"
    ws["I4"] = GLWISE_TITLE
    ws["E8"] = "Date :"
    ws["F8"] = date.strftime("%d-%m-%Y")
    ws["B14"] = "Set ID / Desc :"
    ws["C14"] = "TAMI1 - TAMI1"

    headers = ["S.No", "GL Sub Head Code", "IT2.0 A/C Code", "IT2.0 Acct Code Desc",
               "Deposits (Cr)", "Withdrawals (Dr)"]
    r = 16
    for sol_id, rows in sections:
        ws.cell(row=r, column=2, value="Sol ID / Desc :")
        ws.cell(row=r, column=3, value=f"{sol_id} - Some office")
        r += 1
        for idx, header in enumerate(headers):
            ws.cell(row=r, column=9 + idx, value=header)
        r += 1
        for sno, (code, desc, dep, wd) in enumerate(rows, start=1):
            ws.cell(row=r, column=9, value=sno)
            ws.cell(row=r, column=10, value="30001")
            ws.cell(row=r, column=11, value=int(code))
            ws.cell(row=r, column=12, value=desc)
            ws.cell(row=r, column=13, value=dep)
            ws.cell(row=r, column=14, value=wd)
            r += 1
        ws.cell(row=r, column=12, value="Total")
        r += 2

    wb.save(path)
    return path


def make_apt_details_file(
    path: Path,
    account_code: str,
    rows: list[tuple[dt.date, str, str, float]],  # (date, office_id, office_name, amount)
) -> Path:
    """Build a synthetic APT 'Accounting Details' report for one account code."""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Accounting Details"
    ws["A2"] = f"A/c Code : {account_code}"
    headers = ["Office ID", "Office Name", "Date", "Amount", "Remarks"]
    for c, header in enumerate(headers, start=1):
        ws.cell(row=4, column=c, value=header)
    for r, (date, office_id, office_name, amount) in enumerate(rows, start=5):
        ws.cell(row=r, column=1, value=office_id)
        ws.cell(row=r, column=2, value=office_name)
        ws.cell(row=r, column=3, value=date.strftime("%d/%m/%Y"))
        ws.cell(row=r, column=4, value=amount)
    wb.save(path)
    return path


def make_cashbook_file(
    path: Path,
    date: dt.date,
    rows: list[tuple[str, str, str, float, float, float]],  # (code, desc, side, ho, so, bo)
    office_name: str = "MODEL OFFICE HO",
    office_id: str = "12345600",
) -> Path:
    """Build a synthetic APT Daily Cash Book download: two title rows, then the
    TEMP.CBR column layout."""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Daily Cash Book"
    ws["A2"] = f"Office : {office_name}"
    headers = [
        "Date", "Office Name", "Office ID", "Account Code", "Account Code Description",
        "Part", "Receipts/Payments", "HO", "SO", "BO", "Total", "Progressive Total",
    ]
    for c, header in enumerate(headers, start=1):
        ws.cell(row=3, column=c, value=header)

    progressive = 0.0
    for r, (code, desc, side, ho, so, bo) in enumerate(rows, start=4):
        total = ho + so + bo
        progressive += total
        values = [
            date.strftime("%d/%m/%Y"), office_name, office_id, int(code), desc,
            "Part I", side, ho, so, bo, total, progressive,
        ]
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)

    wb.save(path)
    return path
