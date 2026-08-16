"""Regression tests for the QA and security assessment.

Each test corresponds to a numbered finding. The original suite passed while
all four criticals were live, because it only exercised the happy path and the
legacy defects. These tests exist so those findings cannot come back.
"""

from __future__ import annotations

import csv
import json
import sys
import threading
import urllib.error
import urllib.request
from datetime import date
from decimal import Decimal
from http.server import ThreadingHTTPServer
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from sbco_recon.fiscal import MAX_PERIOD_DAYS, Period
from sbco_recon.ingest import detect_kind, read_rows
from sbco_recon.ingest.batch import load_files, load_one, summarise
from sbco_recon.model import Source
from sbco_recon.normalize import parse_date
from sbco_recon.annexure import build_table1
from sbco_recon.reconcile import (coverage, reconcile_by_code,
                                  reconcile_by_office)
from sbco_recon.reports import write_annexure_iv_table1
from sbco_recon.store import Store

D = Decimal

CASHBOOK_HEADER = ["Date", "Office Name", "Office ID", "Account Code",
                   "Account Code Description", "Part", "Receipts/Payments",
                   "HO", "SO", "BO", "Total", "Progressive Total", "Grand Total"]
FINACLE_HEADER = ["DATE", "SOL ID", "OFFICE", "ACCOUNT CODE", "DESCRIPRTION",
                  "RECEIPT", "VALIDATOR"]
OFFICE_HEADER = ["OFFICE_NAME", "OFFICE_ID", "SOL_ID/BO_CODE", "SOL_ID_GROUP"]


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)
    return path


@pytest.fixture
def store(tmp_path):
    with Store(tmp_path / "qa.db") as st:
        yield st


# ══════════════════════════════════════════════════ QA-01 / QA-02  Critical

def test_qa01_running_total_is_never_read_as_the_amount(tmp_path):
    """A running total placed before the real total used to win on a
    substring match, at score 1.0, inflating a month of figures silently."""
    header = ["Date", "Office Name", "Office ID", "Account Code",
              "Account Code Description", "Part", "Receipts/Payments",
              "Progressive Total", "Total", "Grand Total"]
    path = write_csv(tmp_path / "swapped.csv", header, [
        ["01-07-2026", "HO", "1", "8001000100", "POSB", "I", "R",
         "999999999", "100", "999999999"]])
    d = detect_kind(read_rows(path))
    assert d.kind.value == "cashbook"
    assert header[d.columns["amount"]] == "Total"


@pytest.mark.parametrize("banned", ["Progressive Total", "Grand Total",
                                    "Running Total", "Cumulative Total",
                                    "Closing Balance", "Total To Date"])
def test_qa01_derived_columns_are_disqualified_from_amount(tmp_path, banned):
    header = ["Date", "Account Code", banned, "Total"]
    path = write_csv(tmp_path / "d.csv", header,
                     [["01-07-2026", "8001000100", "999999", "100"]])
    d = detect_kind(read_rows(path))
    if "amount" in d.columns:
        assert header[d.columns["amount"]] != banned


def test_qa02_description_is_never_read_as_the_account_code(tmp_path):
    header = ["Date", "Account Code Description", "Account Code",
              "Total", "Progressive Total"]
    path = write_csv(tmp_path / "desc_first.csv", header,
                     [["01-07-2026", "POSB Receipts", "8001000100", "100", "100"]])
    d = detect_kind(read_rows(path))
    assert header[d.columns["account_code"]] == "Account Code"


def test_qa02_a_convincing_header_over_wrong_data_is_refused(tmp_path):
    """Content verification: the heading says Account Code, the column does
    not hold account codes."""
    path = write_csv(tmp_path / "lying.csv", CASHBOOK_HEADER,
                     [["01-07-2026", "HO", "1", "not-a-code", "POSB", "I", "R",
                       0, 0, 0, "100", "100", "100"]] * 6)
    d = detect_kind(read_rows(path))
    assert not d.ok
    assert "does not hold code values" in d.reason


def test_qa01_totals_are_unchanged_however_the_columns_are_ordered(tmp_path, store):
    """The property test that would have caught both criticals on its own."""
    base = dict(zip(CASHBOOK_HEADER,
                    ["01-07-2026", "HO", "12345600", "8001000100", "POSB",
                     "I", "R", 0, 0, 0, "1234.56", "999999999", "888888888"]))
    # 13! orderings is not a test suite; these four move every field that
    # previously collided - the real total against the running totals, and the
    # account code against its own description.
    seen = set()
    shuffles = [
        CASHBOOK_HEADER,
        list(reversed(CASHBOOK_HEADER)),
        CASHBOOK_HEADER[3:] + CASHBOOK_HEADER[:3],
        [CASHBOOK_HEADER[i] for i in (11, 10, 12, 0, 4, 3, 1, 2, 5, 6, 7, 8, 9)],
    ]
    for n, order in enumerate(shuffles):
        path = write_csv(tmp_path / f"order{n}.csv", order,
                         [[base[h] for h in order]])
        outcome = load_one(store, path)
        assert outcome.ok, f"order {n} rejected: {outcome.reason} {outcome.detail}"
        seen.add(store.totals_by_code(
            Source.CASHBOOK, date(2026, 7, 1), date(2026, 7, 31))["8001000100"])
        store.reverse_batch(outcome.batch_id)
    assert seen == {D("1234.56")}, f"column order changed the figures: {seen}"


# ═══════════════════════════════════════════════════════ QA-03 / QA-07  HTTP

@pytest.fixture
def server(tmp_path):
    from sbco_recon.webapp import server as srv

    srv.Handler.api = srv.Api(tmp_path / "web.db")
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), srv.Handler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    port = httpd.server_address[1]
    yield f"http://127.0.0.1:{port}", srv
    httpd.shutdown()
    httpd.server_close()


def _request(url, method="GET", body=None, headers=None):
    req = urllib.request.Request(url, data=body, method=method,
                                 headers=headers or {})
    try:
        with urllib.request.urlopen(req, timeout=10) as res:
            return res.status, json.loads(res.read() or b"{}")
    except urllib.error.HTTPError as err:
        raw = err.read()
        try:
            return err.code, json.loads(raw or b"{}")
        except json.JSONDecodeError:
            return err.code, {}


def test_qa03_cross_origin_post_is_refused(server):
    """text/plain is CORS-simple, so no preflight protects this route.
    A cross-origin page used to be able to rewrite settings outright."""
    base, srv = server
    status, data = _request(
        f"{base}/api/config", "POST",
        b'{"ddo":"HACKED","ho":"HACKED","division":"HACKED"}',
        {"Origin": "http://evil.example", "Content-Type": "text/plain",
         "X-SBCO-Token": srv.SESSION_TOKEN})
    assert status == 403
    assert "another website" in data["error"]

    _, state = _request(f"{base}/api/state")
    assert state["config"]["ddo"] != "HACKED"


def test_qa03_post_without_the_session_token_is_refused(server):
    base, _ = server
    status, _ = _request(f"{base}/api/config", "POST", b'{"ddo":"X"}',
                         {"Content-Type": "application/json"})
    assert status == 403


def test_qa03_destructive_route_needs_the_token(server):
    base, _ = server
    status, _ = _request(f"{base}/api/reverse", "POST", b'{"batch_id":1}',
                         {"Content-Type": "text/plain"})
    assert status == 403


def test_qa07_foreign_host_header_is_refused(server):
    """Precondition for DNS rebinding."""
    base, _ = server
    status, _ = _request(f"{base}/api/state", headers={"Host": "attacker.example"})
    assert status == 421


def test_qa07_the_database_cannot_be_downloaded(server):
    base, _ = server
    status, data = _request(f"{base}/download?f=..%2F..%2Fsbco_recon.db")
    assert status == 403
    assert "not produced by this session" in data["error"]


def test_qa14_internal_errors_do_not_leak_python_detail(server):
    base, srv = server
    status, data = _request(
        f"{base}/api/transfer-entry", "POST",
        json.dumps({"month": "Jul-2026", "from_code": "1", "to_code": "2",
                    "amount": "not a number"}).encode(),
        {"Content-Type": "application/json", "X-SBCO-Token": srv.SESSION_TOKEN})
    assert "InvalidOperation" not in json.dumps(data)
    assert "ConversionSyntax" not in json.dumps(data)


def test_qa17_lakh_grouped_amounts_are_accepted(server):
    base, srv = server
    status, data = _request(
        f"{base}/api/transfer-entry", "POST",
        json.dumps({"month": "Jul-2026", "from_code": "8001000100",
                    "to_code": "8001000200", "amount": "1,000.50"}).encode(),
        {"Content-Type": "application/json", "X-SBCO-Token": srv.SESSION_TOKEN})
    assert status == 200 and data.get("ok")


def test_qa16_chunked_uploads_are_read_not_silently_truncated(server, tmp_path):
    """A chunked body used to read as zero bytes, so a good file was reported
    empty. It must now arrive intact and be recognised."""
    base, srv = server
    payload = "\n".join(
        [",".join(CASHBOOK_HEADER)]
        + [",".join(["01-07-2026", "HO", "1", "8001000100", "POSB", "I", "R",
                     "0", "0", "0", "100", "100", "100"])] * 3).encode()

    req = urllib.request.Request(
        f"{base}/api/upload", data=iter([payload]), method="POST",
        headers={"X-Filename": "chunked.csv", "Transfer-Encoding": "chunked",
                 "X-SBCO-Token": srv.SESSION_TOKEN})
    with urllib.request.urlopen(req, timeout=10) as res:
        data = json.loads(res.read())
    assert data["status"] == "loaded", data
    assert data["rows"] == 3, f"body truncated: {data}"


def test_qa23_reversing_an_unknown_batch_reports_the_problem(server):
    base, srv = server
    status, data = _request(
        f"{base}/api/reverse", "POST", b'{"batch_id": 9999}',
        {"Content-Type": "application/json", "X-SBCO-Token": srv.SESSION_TOKEN})
    assert "no upload numbered 9999" in data["error"].lower()


# ═════════════════════════════════════════════════════════════ QA-04  Critical

def test_qa04_transfer_entries_reach_the_generated_return(tmp_path, store):
    """The function existed and was unit-tested, but neither call site used
    it, so the return showed uncorrected figures."""
    write_csv(tmp_path / "cb.csv", CASHBOOK_HEADER, [
        ["01-07-2026", "HO", "1", "8001000200", "POSB Pay", "I", "P",
         0, 0, 0, "1000", "", ""]])
    write_csv(tmp_path / "fn.csv", FINACLE_HEADER, [
        ["01-07-2026", "58345610", "HO", "8001000200", "POSB Pay", "1269", ""]])
    load_files(store, [tmp_path / "cb.csv", tmp_path / "fn.csv"])
    store.add_transfer_entry("Jul-2026", "8001000100", "8001000200",
                             D("269"), "approved")

    period = Period(date(2026, 7, 1), date(2026, 7, 31))
    result = reconcile_by_code(store, period, check_coverage=False)
    assert {r.account_code: r.difference
            for r in result.rows}["8001000200"] == D("269")

    # SB Order 09/2026 folds approved TEs into the Monthly Cash Account column
    table1 = build_table1(result, "Jul-2026", store.transfer_entries("Jul-2026"))
    row = {r.account_code: r for r in table1.rows}["8001000200"]
    assert row.difference_payment == 0, "the TE did not reach the figures"

    out = write_annexure_iv_table1(
        tmp_path / "annex.xlsx", table1, ddo_code="102617", ho_name="HO",
        division="D", month=date(2026, 7, 1), include_matched=True)

    ws = load_workbook(out).active
    written = {ws.cell(row=r, column=2).value: r
               for r in range(9, 20) if ws.cell(row=r, column=2).value}
    r = written["8001000200"]
    # E = Finacle payments, G = cash account payments
    assert (Decimal(str(ws.cell(row=r, column=5).value))
            - Decimal(str(ws.cell(row=r, column=7).value))) == 0


def test_qa04_cli_annexure_applies_transfer_entries(tmp_path, monkeypatch):
    from sbco_recon.cli import main
    db = tmp_path / "cli.db"
    write_csv(tmp_path / "cb.csv", CASHBOOK_HEADER, [
        ["01-07-2026", "HO", "1", "8001000200", "POSB Pay", "I", "P",
         0, 0, 0, "1000", "", ""]])
    write_csv(tmp_path / "fn.csv", FINACLE_HEADER, [
        ["01-07-2026", "58345610", "HO", "8001000200", "POSB Pay", "1269", ""]])
    with Store(db) as st:
        load_files(st, [tmp_path / "cb.csv", tmp_path / "fn.csv"])
        st.add_transfer_entry("Jul-2026", "8001000100", "8001000200", D("269"), "")
        st.set("ddo_code", "1"), st.set("ho_name", "HO"), st.set("division", "D")

    out = tmp_path / "cli-annex.xlsx"
    main(["--db", str(db), "annexure", "--month", "Jul-2026",
          "--output", str(out), "--show-all"])
    ws = load_workbook(out).active
    for r in range(9, 30):
        if ws.cell(row=r, column=2).value == "8001000200":
            assert (Decimal(str(ws.cell(row=r, column=5).value))
                    - Decimal(str(ws.cell(row=r, column=7).value))) == 0
            break
    else:
        pytest.fail("8001000200 not written to the return")


# ═════════════════════════════════════════════════════════════════ QA-05

@pytest.mark.parametrize("bad", ["31-02-2026", "31-04-2026", "99-99-2026",
                                 "00-00-2026", "32-01-2026", "01-13-2026"])
def test_qa05_impossible_dates_are_refused_not_rolled_forward(bad):
    assert parse_date(bad) is None


# ═════════════════════════════════════════════════════════════════ QA-06

def test_qa06_closed_days_do_not_raise_a_missing_report_warning(tmp_path, store):
    """A complete month with no Sunday postings used to produce four false
    alarms, training users to ignore the one warning that matters."""
    cb, fn = [], []
    for day in range(1, 32):
        d = date(2026, 7, day)
        if d.weekday() == 6:          # office shut, no rows at all
            continue
        cb.append([d.strftime("%d-%m-%Y"), "HO", "1", "8001000100", "POSB",
                   "I", "R", 0, 0, 0, "1000", "", ""])
        fn.append([d.strftime("%d-%m-%Y"), "58345610", "HO", "8001000100",
                   "POSB", "1000", ""])
    write_csv(tmp_path / "cb.csv", CASHBOOK_HEADER, cb)
    write_csv(tmp_path / "fn.csv", FINACLE_HEADER, fn)
    load_files(store, [tmp_path / "cb.csv", tmp_path / "fn.csv"])

    result = reconcile_by_code(store, Period(date(2026, 7, 1), date(2026, 7, 31)))
    assert not result.differences
    assert not result.warnings, f"false alarm: {result.warnings}"


def test_qa06_a_genuinely_missing_report_still_warns(tmp_path, store):
    """The fix must not silence the real case."""
    write_csv(tmp_path / "cb.csv", CASHBOOK_HEADER, [
        ["01-07-2026", "HO", "1", "8001000100", "POSB", "I", "R",
         0, 0, 0, "1000", "", ""]])
    load_one(store, tmp_path / "cb.csv")
    cov = coverage(store, Period(date(2026, 7, 1), date(2026, 7, 10)))
    assert not cov.complete
    assert len(cov.missing_cashbook) == 9
    assert cov.warnings()


# ═════════════════════════════════════════════════════════════════ QA-08

def test_qa08_exact_figures_survive_the_worksheet(tmp_path, store):
    """The engine carries Decimal so an exact match stays exactly zero;
    float() at the write boundary turned 0.05 into 0.0499267578125."""
    write_csv(tmp_path / "cb.csv", CASHBOOK_HEADER, [
        ["01-07-2026", "HO", "1", "8001000100", "POSB", "I", "R",
         0, 0, 0, "1000000000000.05", "", ""]])
    write_csv(tmp_path / "fn.csv", FINACLE_HEADER, [
        ["01-07-2026", "58345610", "HO", "8001000100", "POSB",
         "1000000000000.10", ""]])
    load_files(store, [tmp_path / "cb.csv", tmp_path / "fn.csv"])
    result = reconcile_by_code(store, Period(date(2026, 7, 1), date(2026, 7, 31)),
                               check_coverage=False)
    assert result.total_difference == D("0.05")

    table1 = build_table1(result, "Jul-2026")
    out = write_annexure_iv_table1(tmp_path / "a.xlsx", table1, ddo_code="1",
                                   ho_name="H", division="D", month=date(2026, 7, 1))
    ws = load_workbook(out).active
    for r in range(9, 20):
        if ws.cell(row=r, column=2).value == "8001000100":
            # D = Finacle receipts, F = cash account receipts
            written = (Decimal(str(ws.cell(row=r, column=4).value))
                       - Decimal(str(ws.cell(row=r, column=6).value)))
            assert written == D("0.05"), f"exported {written}"
            break
    else:
        pytest.fail("row not written")


# ═════════════════════════════════════════════════════════════ QA-10 / QA-11

def test_qa10_replacing_the_office_list_says_what_changed(tmp_path, store):
    write_csv(tmp_path / "a.csv", OFFICE_HEADER, [
        ["HO", "1", "58345610", "58345610"],
        ["SO 1", "2", "58345611", "58345611"]])
    write_csv(tmp_path / "b.csv", OFFICE_HEADER, [
        ["HO", "1", "58345610", "58345610"]])
    assert load_one(store, tmp_path / "a.csv").ok
    outcome = load_one(store, tmp_path / "b.csv")
    assert outcome.ok
    assert "replaced the previous list of 2" in outcome.detail
    assert "1 removed" in outcome.detail


def test_qa11_user_text_cannot_become_a_live_formula(tmp_path, store):
    result = build_table1(
        reconcile_by_code(store, Period(date(2026, 7, 1), date(2026, 7, 31)),
                          check_coverage=False), "Jul-2026")
    out = write_annexure_iv_table1(
        tmp_path / "inject.xlsx", result,
        ddo_code="=1+1", ho_name='=HYPERLINK("http://evil.example","Click")',
        division="D", month=date(2026, 7, 1))
    ws = load_workbook(out).active
    assert not str(ws["B5"].value).startswith("=")
    assert not str(ws["D5"].value).startswith("=")


# ═════════════════════════════════════════════════════════════════ QA-13

def test_qa13_absurd_periods_are_refused_with_a_readable_message():
    with pytest.raises(ValueError) as err:
        Period(date(1900, 1, 1), date(2099, 12, 31))
    assert "at most" in str(err.value)
    Period(date(2026, 4, 1), date(2027, 3, 31))          # a full FY is fine


def test_qa13_reversed_dates_are_refused():
    with pytest.raises(ValueError) as err:
        Period(date(2026, 7, 31), date(2026, 7, 1))
    assert "after the end date" in str(err.value)


# ═════════════════════════════════════════════════════════════ QA-19 / QA-21

def test_qa19_duplicate_is_reported_not_raised_as_an_integrity_error(tmp_path, store):
    """Insert-then-catch, so a second uploader gets 'duplicate' rather than a
    500 from the partial unique index."""
    path = write_csv(tmp_path / "cb.csv", CASHBOOK_HEADER, [
        ["01-07-2026", "HO", "1", "8001000100", "POSB", "I", "R",
         0, 0, 0, "100", "", ""]])
    first = load_one(store, path)
    assert first.ok
    # bypass the pre-check the way a racing process would
    from sbco_recon.ingest.parse import parse_entries
    rows = read_rows(path)
    entries, _ = parse_entries(rows, detect_kind(rows), path)
    with pytest.raises(Store.Duplicate):
        store.add_batch(Source.CASHBOOK, path, first.sha256, entries)


def test_qa21_the_accounting_invariant_is_not_an_assert(tmp_path, store):
    """`python -O` strips assertions; this guarantee must survive it."""
    import sbco_recon.ingest.batch as batch_mod
    source = Path(batch_mod.__file__).read_text(encoding="utf-8")
    assert "assert " not in source, "an assert crept back into the batch loader"

    counts = summarise(load_files(store, [tmp_path / "nope1.xls",
                                          tmp_path / "nope2.xls"]))
    assert counts["submitted"] == 2
    assert counts["rejected"] == 2          # "file not found", each named
    assert (counts["loaded"] + counts["rejected"]
            + counts["duplicate"] + counts["failed"]) == 2


# ═════════════════════════════════════════════════════════════════ QA-25

def test_qa25_office_matching_does_not_cross_id_namespaces(tmp_path, store):
    """A BO code equal to another office's ID used to move money silently."""
    write_csv(tmp_path / "off.csv", OFFICE_HEADER, [
        ["HO", "58345611", "58345610", "58345610"],   # office_id == SO's sol_id
        ["SO 1", "12345601", "58345611", "58345611"]])
    load_one(store, tmp_path / "off.csv")
    write_csv(tmp_path / "gl.csv", FINACLE_HEADER, [
        ["01-07-2026", "58345611", "SO", "8001000100", "POSB", "500", ""]])
    load_one(store, tmp_path / "gl.csv")

    result = reconcile_by_office(store, "8001000100",
                                 Period(date(2026, 7, 1), date(2026, 7, 31)))
    by_name = {r.office.name: r for r in result.rows}
    assert by_name["SO 1"].finacle == D("500")     # matched on SOL ID
    assert by_name["HO"].finacle == D("0")         # not on its office ID


# ═════════════════════════════════════════════ v2.0 -> v2.1 data migration

def test_upgrade_carries_the_whole_database_across(tmp_path, monkeypatch):
    """Before 2.1 the launcher created the database in the program folder.
    Upgrading must not look like total data loss.

    The copy goes through SQLite's backup API: the database runs in WAL mode,
    so a plain file copy silently drops the most recent commits - it carried
    the settings over but none of the loaded reports, which is worse than
    failing outright.
    """
    from sbco_recon.store import adopt_legacy_database

    legacy = tmp_path / "program" / "sbco_recon.db"
    legacy.parent.mkdir()
    write_csv(tmp_path / "cb.csv", CASHBOOK_HEADER, [
        ["01-07-2026", "HO", "1", "8001000100", "POSB", "I", "R",
         0, 0, 0, "1000", "", ""]])
    write_csv(tmp_path / "off.csv", OFFICE_HEADER, [
        ["HO", "12345600", "58345610", "58345610"]])

    with Store(legacy) as old:
        load_files(old, [tmp_path / "cb.csv", tmp_path / "off.csv"])
        old.set("ddo_code", "102617")
        before = reconcile_by_code(old, Period(date(2026, 7, 1), date(2026, 7, 31)),
                                   check_coverage=False).total_difference
        batches_before = len(old.batches())

    target = tmp_path / "userdata" / "sbco_recon.db"
    target.parent.mkdir()
    monkeypatch.chdir(legacy.parent)
    assert adopt_legacy_database(target) is True

    with Store(target) as new:
        assert new.get("ddo_code") == "102617"
        assert len(new.offices()) == 1
        assert len(new.batches()) == batches_before, "loaded reports were dropped"
        after = reconcile_by_code(new, Period(date(2026, 7, 1), date(2026, 7, 31)),
                                  check_coverage=False).total_difference
    assert after == before

    assert legacy.exists(), "the original must survive, so 2.0 can be rolled back"
    assert adopt_legacy_database(target) is False, "must not run twice"


def test_computing_the_default_path_does_not_migrate_anything(tmp_path, monkeypatch):
    """The migration once ran as a side effect of building the argparse
    default, firing during unrelated commands and copying a half-written
    database."""
    from sbco_recon.store import default_db_path

    monkeypatch.setenv("SBCO_DATA_DIR", str(tmp_path / "data"))
    legacy = tmp_path / "sbco_recon.db"
    with Store(legacy) as old:
        old.set("ddo_code", "SHOULD-NOT-TRAVEL")
    monkeypatch.chdir(tmp_path)

    target = default_db_path()
    assert not target.exists(), "computing a path must not create a database"


# ═══════════════════════════════════════════════ version number consistency

def test_one_version_number_everywhere():
    """The legacy tool showed 1.04 on screen while its HELP sheet said 1.09.6.
    Shipping an archive named 2.1.0 containing __version__ = "2.0.0" is the
    same defect. When a user reports a bug, the version they quote must be
    the version they are running.
    """
    import tomllib

    from sbco_recon import __version__

    root = Path(__file__).resolve().parents[1]
    with open(root / "pyproject.toml", "rb") as fh:
        declared = tomllib.load(fh)["project"]["version"]
    assert declared == __version__, (
        f"pyproject.toml says {declared}, __init__.py says {__version__}")


def test_doctor_never_changes_anything(tmp_path, monkeypatch, capsys):
    """A diagnostic must be read-only. `doctor` used to trigger the pre-2.1
    data migration, consuming the very file it was meant to report on."""
    from sbco_recon.cli import main

    monkeypatch.setenv("SBCO_DATA_DIR", str(tmp_path / "data"))
    legacy = tmp_path / "sbco_recon.db"
    with Store(legacy) as old:
        old.set("ddo_code", "102617")
    monkeypatch.chdir(tmp_path)

    main(["doctor"])
    out = capsys.readouterr().out

    from sbco_recon.store import default_db_path
    assert not default_db_path().exists(), "doctor created a data file"
    assert "older version was found" in out, "doctor should report it, not adopt it"


def test_a_real_xlsx_named_dot_xls_is_read_not_refused(tmp_path):
    """Portals routinely emit an OOXML workbook called REPORT.xls. openpyxl
    decides what a file is from its extension and refuses anything ending
    .xls, so passing the path rejected a perfectly readable file with an
    error about 'the old .xls format'. The content is what matters."""
    from openpyxl import Workbook

    from sbco_recon.ingest import detect_kind, read_rows

    wb = Workbook()
    ws = wb.active
    ws.append(CASHBOOK_HEADER)
    for _ in range(3):
        ws.append(["01-07-2026", "HO", "1", "8001000100", "POSB", "I", "R",
                   0, 0, 0, "100", "100", "100"])
    misnamed = tmp_path / "REPORT.xls"
    wb.save(misnamed)
    assert misnamed.read_bytes()[:2] == b"PK", "fixture should be a zip"

    rows = read_rows(misnamed)
    assert len(rows) == 4
    assert detect_kind(rows).kind.value == "cashbook"


def test_an_html_table_named_dot_xls_is_read(tmp_path):
    """The other thing portals emit under an .xls name."""
    from sbco_recon.ingest import detect_kind, read_rows

    cells = "".join(f"<td>{h}</td>" for h in CASHBOOK_HEADER)
    data = "".join(f"<td>{v}</td>" for v in
                   ["01-07-2026", "HO", "1", "8001000100", "POSB", "I", "R",
                    0, 0, 0, "100", "100", "100"])
    path = tmp_path / "REPORT.xls"
    path.write_text(f"<html><body><table><tr>{cells}</tr>"
                    f"<tr>{data}</tr></table></body></html>")
    assert detect_kind(read_rows(path)).kind.value == "cashbook"
