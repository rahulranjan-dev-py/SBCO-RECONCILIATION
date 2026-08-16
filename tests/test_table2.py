"""Annexure-IV Table-2, against the layout recovered from v1.09.8.

The arithmetic mirrors the original's own worksheet formulas:
    closing = opening + current - rectified, per side
with the side taken from the ac_codes master, not from the sign.
"""

from __future__ import annotations

import sys
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from sbco_recon.table2 import (PAYMENT, RECEIPT, Table2Row, build_table2,
                               months_between, side_of)

D = Decimal

RECEIPT_CODE = "8001000100"    # POSB Account -Receipts
PAYMENT_CODE = "8001000200"    # POSB Account -Payments


@dataclass
class Row:
    account_code: str
    description: str
    difference: Decimal


# ── side classification ──────────────────────────────────────────────

def test_side_comes_from_the_account_code_master():
    assert side_of(RECEIPT_CODE) == RECEIPT
    assert side_of(PAYMENT_CODE) == PAYMENT
    assert side_of("8671001900") == PAYMENT          # HR to CBS Cleared


def test_side_matching_is_case_insensitive():
    """The recovered master holds one code as 'payment Side' with a small p.
    Excel's = ignores case; Python's does not, so a literal match would have
    filed that code on the wrong side of a statutory return."""
    from sbco_recon import refdata
    odd = [c for c, v in refdata.account_codes().items()
           if (v.get("side") or "") not in ("Receipt Side", "Payment Side")
           and v.get("side")]
    assert odd, "expected the known case anomaly in the master"
    for code in odd:
        assert side_of(code) in (RECEIPT, PAYMENT)
        assert side_of(code) == PAYMENT


def test_an_unknown_code_still_gets_a_side():
    assert side_of("9999999999") in (RECEIPT, PAYMENT)


# ── the carry-forward arithmetic ─────────────────────────────────────

def test_a_receipt_difference_lands_in_the_receipts_column():
    result = build_table2("Jul-2026", [Row(RECEIPT_CODE, "POSB Rec", D("1500"))])
    row = result.rows[0]
    assert row.current_receipt == D("1500")
    assert row.current_payment == 0
    assert row.closing_receipt == D("1500")


def test_a_payment_difference_lands_in_the_payments_column():
    result = build_table2("Jul-2026", [Row(PAYMENT_CODE, "POSB Pay", D("24"))])
    row = result.rows[0]
    assert row.current_payment == D("24")
    assert row.current_receipt == 0


def test_side_is_decided_by_the_code_not_the_sign():
    """A negative difference on a receipt code is still a receipt."""
    result = build_table2("Jul-2026", [Row(RECEIPT_CODE, "POSB Rec", D("-900"))])
    row = result.rows[0]
    assert row.current_receipt == D("-900")
    assert row.current_payment == 0


def test_opening_is_taken_from_the_previous_months_closing():
    """Not from its opening - the whole point of the roll."""
    prior = Table2Row(PAYMENT_CODE, "POSB Pay", opening_payment=D("100"))
    prior.current_payment = D("40")
    prior.rectified_payment = D("30")
    assert prior.closing_payment == D("110")     # 100 + 40 - 30

    july = build_table2("Jul-2026", [], previous_table2=[prior])
    assert july.rows[0].opening_payment == D("110")
    assert july.rows[0].closing_payment == D("110")
    assert july.balances


def test_closing_equals_opening_plus_current_less_rectified():
    prior = Table2Row(PAYMENT_CODE, "POSB Pay", opening_payment=D("500"))
    row = build_table2(
        "Jul-2026", [Row(PAYMENT_CODE, "POSB Pay", D("269"))],
        previous_table2=[prior],
        transfer_entries=[(PAYMENT_CODE, D("200"))]).rows[0]
    assert row.opening_payment == D("500")
    assert row.current_payment == D("269")
    assert row.rectified_payment == D("200")
    assert row.closing_payment == D("569")       # 500 + 269 - 200


def test_a_transfer_entry_rectifies_on_the_codes_own_side():
    seed = Table2Row(PAYMENT_CODE, "POSB Pay")
    seed.current_payment = D("269")
    result = build_table2("Aug-2026", [], previous_table2=[seed],
                          transfer_entries=[(PAYMENT_CODE, D("269"))])
    row = result.rows[0]
    assert row.rectified_payment == D("269")
    assert row.rectified_receipt == 0
    assert row.closing_payment == 0
    assert row.is_settled
    assert result.balances


def test_items_persist_until_rectified():
    jul = build_table2("Jul-2026", [Row(PAYMENT_CODE, "POSB Pay", D("269"))])
    aug = build_table2("Aug-2026", [], previous_table2=jul.rows)
    sep = build_table2("Sep-2026", [], previous_table2=aug.rows)
    assert sep.closing_total == D("269")
    assert sep.rows[0].first_month == "Jul-2026", "age must survive the roll"

    oct_ = build_table2("Oct-2026", [], previous_table2=sep.rows,
                        transfer_entries=[(PAYMENT_CODE, D("269"))])
    assert oct_.closing_total == 0
    assert oct_.balances


def test_the_return_always_ties_back():
    seed = Table2Row("8671001900", "HR to CBS")
    seed.current_payment = D("-427025")
    result = build_table2(
        "Jul-2026",
        [Row(RECEIPT_CODE, "POSB Rec", D("1500")),
         Row(PAYMENT_CODE, "POSB Pay", D("24"))],
        previous_table2=[seed],
        transfer_entries=[("8671001900", D("-27025"))])
    assert result.balances
    assert (result.opening_total + result.current_total
            - result.rectified_total) == result.closing_total


def test_receipts_and_payments_never_bleed_into_each_other():
    result = build_table2("Jul-2026", [
        Row(RECEIPT_CODE, "POSB Rec", D("500")),
        Row(PAYMENT_CODE, "POSB Pay", D("300"))])
    by_code = {r.account_code: r for r in result.rows}
    assert by_code[RECEIPT_CODE].closing_payment == 0
    assert by_code[PAYMENT_CODE].closing_receipt == 0


def test_a_transfer_entry_for_an_unknown_code_is_flagged():
    result = build_table2("Jul-2026", [],
                          transfer_entries=[("8446014200", D("100"))])
    assert any("nothing outstanding" in w for w in result.warnings)
    assert result.balances


def test_settled_codes_still_appear_on_the_return():
    """A code opened and cleared in the same month must be visible, not
    silently dropped - the PAO needs to see it was rectified."""
    result = build_table2("Jul-2026", [Row(PAYMENT_CODE, "POSB Pay", D("269"))],
                          transfer_entries=[(PAYMENT_CODE, D("269"))])
    assert len(result.rows) == 1
    assert result.rows[0].is_settled
    assert result.rows[0].rectified_payment == D("269")


def test_codes_with_no_activity_are_left_off():
    result = build_table2("Jul-2026", [Row(RECEIPT_CODE, "x", D("0"))])
    assert not result.rows


def test_transfer_entries_accept_dicts_and_rows():
    seed = Table2Row(PAYMENT_CODE, "POSB Pay")
    seed.current_payment = D("100")
    result = build_table2("Aug-2026", [], previous_table2=[seed],
                          transfer_entries=[{"account_code": PAYMENT_CODE,
                                             "amount": "100"}])
    assert result.rows[0].closing_payment == 0


def test_ageing_groups_pending_items(monkeypatch):
    rows = []
    for code, first in ((RECEIPT_CODE, "Jul-2026"), (PAYMENT_CODE, "Apr-2026"),
                        ("8671001900", "Jan-2026"), ("8446014200", "Jun-2025")):
        r = Table2Row(code, "", first_month=first)
        r.current_receipt = D("10") if side_of(code) == RECEIPT else D("0")
        r.current_payment = D("10") if side_of(code) == PAYMENT else D("0")
        rows.append(r)
    result = build_table2("Jul-2026", [], previous_table2=rows)
    for r, first in zip(result.rows, sorted(rows, key=lambda x: x.account_code)):
        r.first_month = first.first_month
    buckets = result.ageing()
    assert sum(len(v) for v in buckets.values()) == 4


@pytest.mark.parametrize("a,b,expected", [
    ("Jul-2026", "Jul-2026", 0), ("Jun-2026", "Jul-2026", 1),
    ("Dec-2025", "Mar-2026", 3), ("Apr-2026", "Mar-2027", 11),
    ("nonsense", "Jul-2026", 0),
])
def test_month_arithmetic_crosses_year_ends(a, b, expected):
    assert months_between(a, b) == expected


# ── the exported workbook ────────────────────────────────────────────

def test_the_exported_return_matches_the_recovered_layout(tmp_path):
    """Column positions copied from v1.09.8 so the PAO gets the expected form."""
    from datetime import date as _date

    from openpyxl import load_workbook

    from sbco_recon.reports import write_annexure_iv_table2

    seed = Table2Row("8671001900", "HR to CBS Cleared")
    seed.current_payment = D("-427025")
    result = build_table2("Aug-2026",
                          [Row(RECEIPT_CODE, "POSB Rec", D("1500"))],
                          previous_table2=[seed])
    out = write_annexure_iv_table2(tmp_path / "t2.xlsx", result, ddo_code="102617",
                                   ho_name="Manipal HO", division="Manipal",
                                   month=_date(2026, 8, 1))
    ws = load_workbook(out).active

    assert ws["A1"].value.startswith("Detailed CBS Monthly Reconciliation")
    assert ws["A4"].value == "Table-2"
    assert ws["A7"].value == "SL NO."
    assert ws["B7"].value == "A/c Code"
    assert ws["C7"].value == "A/c Code Description"
    assert "Opening" in ws["D7"].value and "Current" in ws["F7"].value
    assert "Rectified" in ws["H7"].value
    assert ws["D8"].value.startswith("Receipts")
    assert ws["E8"].value.startswith("Payments")

    body = {ws.cell(row=r, column=2).value: r for r in range(9, 20)
            if ws.cell(row=r, column=2).value}
    r = body[RECEIPT_CODE]
    assert ws.cell(row=r, column=6).value == D("1500")     # F: current receipts
    assert ws.cell(row=r, column=7).value == 0             # G: current payments
    assert ws.cell(row=r, column=10).value == f"=D{r}+F{r}-H{r}"


def test_exported_figures_stay_exact(tmp_path):
    from datetime import date as _date

    from openpyxl import load_workbook

    from sbco_recon.reports import write_annexure_iv_table2

    result = build_table2("Jul-2026",
                          [Row(PAYMENT_CODE, "POSB Pay", D("1000000000000.05"))])
    out = write_annexure_iv_table2(tmp_path / "t2.xlsx", result, ddo_code="1",
                                   ho_name="H", division="D", month=_date(2026, 7, 1))
    ws = load_workbook(out).active
    for r in range(9, 15):
        if ws.cell(row=r, column=2).value == PAYMENT_CODE:
            assert D(str(ws.cell(row=r, column=7).value)) == D("1000000000000.05")
            break
    else:
        pytest.fail("row not written")
