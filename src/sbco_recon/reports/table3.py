"""Table-3: CBS Daily Discrepancy Reconciliation Register.

Prescribed by SB Order No. 09/2026, Annexure-IV. Para 4(vi) requires SBCO to
maintain it for every discrepancy detected, and to preserve it permanently;
para 7 has inspecting authorities check it during HO inspection.

Columns (a)-(o) follow the order's own labelling. The two difference columns
are formulas, because the order defines them as (f)-(h) and (g)-(i) - written
as values they could drift from the figures beside them.

The register closes at the end of each financial year and the serial numbering
restarts from 1, which is why entries carry a financial year.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ..model import ZERO
from . import style

TITLE = "CBS Daily Discrepancy Reconciliation Register"
FOOTNOTE = ("* The Register shall be closed at the end of each financial year, "
            "and the serial numbering shall reset from S No.1 at the beginning "
            "of the next financial year.")

# (label, order-letter, column width)
COLUMNS = (
    ("Sl. No.", "a", 7),
    ("Date", "b", 12),
    ("Account Code", "c", 14),
    ("Description of Account Code", "d", 34),
    ("Name of the Post office in which the discrepancy is found", "e", 26),
    ("Receipt as per CBS\n(in Rs.)", "f", 15),
    ("Payment as per CBS\n(in Rs.)", "g", 15),
    ("Receipt as per Cash Book\n(in Rs.)", "h", 15),
    ("Payment as per Cash Book\n(in Rs.)", "i", 15),
    ("Difference (CBS - Cashbook)\nin Receipt (in Rs.)", "(f) - (h)", 16),
    ("Difference (CBS - Cashbook)\nin Payment (in Rs.)", "(g) - (i)", 16),
    ("Initials of In-Charge SBCO", "j", 13),
    ("Date of Rectification", "k", 14),
    ("Particulars of Misc. Transaction Posted", "l", 24),
    ("Particulars of Transfer Entries Posted", "m", 24),
    ("Initials of PA / APM", "n", 12),
    ("Initials of Postmaster", "o", 12),
)

HEAD_ROW, LETTER_ROW, FIRST_DATA_ROW = 5, 6, 7
RECEIPT_DIFF_COL, PAYMENT_DIFF_COL = 10, 11


def write_discrepancy_register(path, entries, *, financial_year: str = "",
                               ho_name: str = "") -> Path:
    """Write the register for one financial year."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Table-3"
    last_col = len(COLUMNS)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(row=1, column=1, value=TITLE).font = style.TITLE
    ws.cell(row=1, column=1).alignment = style.CENTER

    if financial_year or ho_name:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        parts = [p for p in (style.safe_text(ho_name),
                             f"Financial Year {financial_year}" if financial_year else "")
                 if p]
        ws.cell(row=2, column=1, value="   |   ".join(parts)).font = style.SUBTITLE
        ws.cell(row=2, column=1).alignment = style.CENTER

    for idx, (label, letter, width) in enumerate(COLUMNS, start=1):
        head = ws.cell(row=HEAD_ROW, column=idx, value=label)
        style.apply_header(head)
        tag = ws.cell(row=LETTER_ROW, column=idx, value=letter)
        tag.font = style.SUBTITLE
        tag.alignment = style.CENTER
        tag.border = style.BOX
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[HEAD_ROW].height = 62

    for n, entry in enumerate(entries):
        r = FIRST_DATA_ROW + n
        ws.cell(row=r, column=1,
                value=entry.serial or (n + 1)).alignment = style.CENTER
        date_cell = ws.cell(row=r, column=2, value=entry.entry_date)
        date_cell.number_format = style.DATE_FMT
        style.write_text(ws, r, 3, entry.account_code)
        style.write_text(ws, r, 4, entry.description)
        style.write_text(ws, r, 5, entry.office_name)

        for col, value in ((6, entry.cbs_receipt), (7, entry.cbs_payment),
                           (8, entry.cashbook_receipt), (9, entry.cashbook_payment)):
            style.write_money(ws, r, col, value).font = style.BODY

        for col, cbs, book, amount in (
                (RECEIPT_DIFF_COL, "F", "H", entry.difference_receipt),
                (PAYMENT_DIFF_COL, "G", "I", entry.difference_payment)):
            cell = ws.cell(row=r, column=col, value=f"={cbs}{r}-{book}{r}")
            cell.font = style.NEGATIVE if amount < ZERO else style.BODY
            cell.number_format = style.CURRENCY
            cell.alignment = style.RIGHT

        style.write_text(ws, r, 12, entry.sbco_initials)
        if entry.rectified_date:
            rect = ws.cell(row=r, column=13, value=entry.rectified_date)
            rect.number_format = style.DATE_FMT
        else:
            ws.cell(row=r, column=13, value="")
        style.write_text(ws, r, 14, entry.misc_transaction)
        style.write_text(ws, r, 15, entry.transfer_entry)
        style.write_text(ws, r, 16, entry.pa_initials)
        style.write_text(ws, r, 17, entry.postmaster_initials)

        for col in range(1, last_col + 1):
            ws.cell(row=r, column=col).border = style.BOX

    note = FIRST_DATA_ROW + len(list(entries)) + 2
    ws.merge_cells(start_row=note, start_column=1, end_row=note, end_column=last_col)
    ws.cell(row=note, column=1, value=FOOTNOTE).font = style.SUBTITLE

    ws.cell(row=note + 3, column=1, value="SBCO In-charge").font = style.BOLD
    ws.cell(row=note + 4, column=1, value="HO").font = style.BODY

    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=3)
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = f"{HEAD_ROW}:{LETTER_ROW}"
    wb.save(Path(path))
    return Path(path)
