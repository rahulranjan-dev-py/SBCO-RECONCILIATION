"""Annexure-IV Table-2 workbook writer.

Column layout copied from CASHBOOK TOOL for SBCO 1.09.8 so the return the PAO
receives is the one they expect:

    A  SL NO.        D/E  Opening Balance      receipts / payments
    B  A/c Code      F/G  Current Month        receipts / payments
    C  Description   H/I  Rectified            receipts / payments
                     J/K  Closing              receipts / payments

The original also carries helper columns L-R (TE totals, side lookups and
VLOOKUP scratch). Those are internals of the Excel implementation, not part of
the prescribed return, so they are not written here.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ..model import ZERO
from . import style

TITLE = "Detailed CBS Monthly Reconciliation Report to PAO by the HO"
GROUPS = (("Opening Balance in the Difference\n(Finacle - Cash Account)", "D", "E"),
          ("Current Month Difference\n(Finacle - Cash Account)", "F", "G"),
          ("Rectified During the Current Month\n(Finacle - Cash Account)", "H", "I"),
          ("Pending for Rectification", "J", "K"))

HEAD_ROW = 7
SUB_ROW = 8
FIRST_DATA_ROW = 9


def write_annexure_iv_table2(path, result, *, ddo_code: str, ho_name: str,
                             division: str, month: date,
                             include_settled: bool = True) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Table-2"

    ws.merge_cells("A1:K1")
    ws["A1"] = TITLE
    ws["A1"].font = style.TITLE
    ws["A1"].alignment = style.CENTER

    ws.merge_cells("A2:K2")
    ws["A2"] = "(Due Date: 4th of every month)"
    ws["A2"].font = style.SUBTITLE
    ws["A2"].alignment = style.CENTER

    ws.merge_cells("A4:K4")
    ws["A4"] = "Table-2"
    ws["A4"].font = style.BOLD
    ws["A4"].alignment = style.CENTER

    ws["A6"] = "DDO Code:"
    ws["A6"].font = style.BOLD
    ws["B6"] = style.safe_text(ddo_code)
    ws["D6"] = "HO:"
    ws["D6"].font = style.BOLD
    ws["E6"] = style.safe_text(ho_name)
    ws["G6"] = "Division:"
    ws["G6"].font = style.BOLD
    ws["H6"] = style.safe_text(division)
    ws["J6"] = "Month:"
    ws["J6"].font = style.BOLD
    ws["K6"] = month.strftime("%b-%Y")

    for ref, label in (("A", "SL NO."), ("B", "A/c Code"),
                       ("C", "A/c Code Description")):
        ws.merge_cells(f"{ref}{HEAD_ROW}:{ref}{SUB_ROW}")
        style.apply_header(ws[f"{ref}{HEAD_ROW}"])
        ws[f"{ref}{HEAD_ROW}"] = label
        style.apply_header(ws[f"{ref}{SUB_ROW}"])

    for label, left, right in GROUPS:
        ws.merge_cells(f"{left}{HEAD_ROW}:{right}{HEAD_ROW}")
        ws[f"{left}{HEAD_ROW}"] = label
        style.apply_header(ws[f"{left}{HEAD_ROW}"])
        style.apply_header(ws[f"{right}{HEAD_ROW}"])
        for ref, side in ((left, "Receipts\n(In Rs.)"), (right, "Payments\n(In Rs.)")):
            ws[f"{ref}{SUB_ROW}"] = side
            style.apply_header(ws[f"{ref}{SUB_ROW}"])

    ws.row_dimensions[HEAD_ROW].height = 46
    ws.row_dimensions[SUB_ROW].height = 30

    rows = result.rows if include_settled else result.pending
    rows = sorted(rows, key=lambda r: r.account_code)

    for n, item in enumerate(rows, start=1):
        r = FIRST_DATA_ROW + n - 1
        ws.cell(row=r, column=1, value=n).alignment = style.CENTER
        style.write_text(ws, r, 2, item.account_code)
        style.write_text(ws, r, 3, item.description)

        for col, value in ((4, item.opening_receipt), (5, item.opening_payment),
                           (6, item.current_receipt), (7, item.current_payment),
                           (8, item.rectified_receipt), (9, item.rectified_payment)):
            cell = style.write_money(ws, r, col, value)
            cell.font = style.BODY

        # closing stays a live formula, so the PAO can see how it was derived
        for col, letter, opening, current, rectified in (
                (10, "J", "D", "F", "H"), (11, "K", "E", "G", "I")):
            cell = ws.cell(row=r, column=col,
                           value=f"={opening}{r}+{current}{r}-{rectified}{r}")
            closing = (item.closing_receipt if col == 10 else item.closing_payment)
            cell.font = style.NEGATIVE if closing < ZERO else style.BODY
            cell.number_format = style.CURRENCY
            cell.alignment = style.RIGHT

        for col in range(1, 12):
            ws.cell(row=r, column=col).border = style.BOX

    last = FIRST_DATA_ROW + len(rows) - 1
    total_row = FIRST_DATA_ROW + len(rows)
    ws.cell(row=total_row, column=3, value="TOTAL").font = style.BOLD
    for col in range(4, 12):
        letter = get_column_letter(col)
        cell = ws.cell(row=total_row, column=col,
                       value=(f"=SUM({letter}{FIRST_DATA_ROW}:{letter}{last})"
                              if rows else 0))
        cell.font = style.BOLD
        cell.number_format = style.CURRENCY
        cell.alignment = style.RIGHT
        cell.fill = style.TOTAL_FILL
        cell.border = style.BOX
    for col in (1, 2, 3):
        c = ws.cell(row=total_row, column=col)
        c.fill = style.TOTAL_FILL
        c.border = style.BOX

    # The order requires both of these under the table; they are left blank
    # for SBCO to complete rather than invented.
    note = total_row + 2
    ws.cell(row=note, column=1,
            value="(a) Reasons for pending for rectification :").font = style.BODY
    ws.cell(row=note + 1, column=1,
            value="(b) Date by which the pendency is cleared :").font = style.BODY
    for row in (note, note + 1):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        for col in range(5, 12):
            ws.cell(row=row, column=col).border = style.BOX

    extra = note + 3
    ws.cell(row=extra, column=1,
            value="Pending for Rectification = Opening + Current month - "
                  "Rectified, for each of receipts and payments.").font = style.SUBTITLE
    for i, warning in enumerate(result.warnings, start=1):
        ws.cell(row=extra + i, column=1, value=warning).font = style.SUBTITLE

    sign = extra + len(result.warnings) + 3
    ws.cell(row=sign, column=2, value="SBCO In-charge").font = style.BOLD
    ws.cell(row=sign, column=8, value="Postmaster").font = style.BOLD
    ws.cell(row=sign + 1, column=2, value="HO").font = style.BODY
    ws.cell(row=sign + 1, column=8, value="HO").font = style.BODY
    ws.cell(row=sign + 3, column=1,
            value="Forwarded to the General Manager(F) / DA(P)").font = style.BODY
    ws.cell(row=sign + 4, column=1,
            value=style.safe_text(
                f"Copy to: The S/SPOs, {division} Division for information."
            )).font = style.BODY

    style.autosize(ws, [8, 14, 44, 15, 15, 15, 15, 15, 15, 15, 15])
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = f"{HEAD_ROW}:{SUB_ROW}"

    out = Path(path)
    wb.save(out)
    return out
