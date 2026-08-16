"""General-purpose Excel exports: reconciliation sheets and the discrepancy register."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from ..model import ZERO
from . import style


def write_recon_sheet(path, result, *, title="Account Code Reconciliation",
                      include_matched=True) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reconciliation"

    ws.merge_cells("A1:E1")
    ws["A1"] = title
    ws["A1"].font = style.TITLE
    ws["A1"].alignment = style.CENTER
    ws["A2"] = f"Period: {result.period}"
    ws["A2"].font = style.SUBTITLE

    headers = ["A/c Code", "Description", "Finacle / CBS", "Cash Book", "Difference"]
    for col, text in enumerate(headers, start=1):
        style.apply_header(ws.cell(row=4, column=col, value=text))

    rows = result.rows if include_matched else result.differences
    first = 5
    for n, row in enumerate(rows):
        r = first + n
        style.write_text(ws, r, 1, row.account_code)
        style.write_text(ws, r, 2, row.description)
        style.write_money(ws, r, 3, row.finacle).font = style.BODY
        style.write_money(ws, r, 4, row.cashbook).font = style.BODY
        diff = ws.cell(row=r, column=5, value=f"=C{r}-D{r}")
        diff.font = style.NEGATIVE if row.difference < ZERO else style.BODY
        for col in range(1, 6):
            cell = ws.cell(row=r, column=col)
            cell.border = style.BOX
            if col >= 3:
                cell.number_format = style.CURRENCY

    last = first + len(rows) - 1
    total = first + len(rows)
    ws.cell(row=total, column=2, value="TOTAL").font = style.BOLD
    for col, letter in ((3, "C"), (4, "D"), (5, "E")):
        cell = ws.cell(row=total, column=col,
                       value=f"=SUM({letter}{first}:{letter}{last})" if rows else 0)
        cell.font = style.BOLD
        cell.number_format = style.CURRENCY
        cell.fill = style.TOTAL_FILL
        cell.border = style.BOX

    if result.warnings:
        wrow = total + 2
        ws.cell(row=wrow, column=1, value="Warnings").font = style.BOLD
        for n, warning in enumerate(result.warnings, start=1):
            ws.cell(row=wrow + n, column=1, value=warning).font = style.SUBTITLE

    style.autosize(ws, [14, 52, 18, 18, 18])
    ws.freeze_panes = "A5"
    wb.save(Path(path))
    return Path(path)


def write_discrepancy_report(path, rows) -> Path:
    """Export the saved discrepancy register.

    Unlike the legacy DESC_RPT_EXPORT, this never refuses to export because a
    column happens to sum to zero - that check used Or where it needed And.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Discrepancy Report"

    ws.merge_cells("A1:H1")
    ws["A1"] = "DISCREPANCY REPORT"
    ws["A1"].font = style.TITLE
    ws["A1"].alignment = style.CENTER

    headers = ["FROM", "TO", "A/C CODE", "DESCRIPTION", "CBS DATA",
               "APT DATA", "DIFFERENCE", "REMARKS"]
    for col, text in enumerate(headers, start=1):
        style.apply_header(ws.cell(row=3, column=col, value=text))

    for n, row in enumerate(rows, start=4):
        for col, value in enumerate(
                [row["period_start"], row["period_end"], row["account_code"],
                 row["description"], row["remarks"]], start=1):
            col = col if col < 5 else 8
            style.write_text(ws, n, col, value).border = style.BOX
        for col, key in ((5, "finacle"), (6, "cashbook"), (7, "difference")):
            cell = style.write_money(ws, n, col, row[key])
            cell.font = style.BODY
            cell.border = style.BOX

    style.autosize(ws, [13, 13, 14, 46, 16, 16, 16, 34])
    ws.freeze_panes = "A4"
    wb.save(Path(path))
    return Path(path)
