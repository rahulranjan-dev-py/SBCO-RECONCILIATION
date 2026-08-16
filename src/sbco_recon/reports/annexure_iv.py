"""Annexure-IV Table-1: CBS Monthly Reconciliation Report to PAO by the HO.

Layout exactly as prescribed by SB Order No. 09/2026, Annexure-IV:

    Sl | Account Code | Account Code Description
       | Finacle              Receipts | Payments
       | Monthly Cash Account*  Receipts | Payments
       | Difference             Receipts | Payments

Every money column is split by side. Para 1(ix) of the order forbids netting a
receipt discrepancy against a payment one - a single signed column per code
would do exactly that.

Signed jointly by the In-charge SBCO and the Postmaster, per para 4(ix).
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ..annexure import CASH_ACCOUNT_NOTE
from ..model import ZERO
from . import style

TITLE = "CBS Monthly Reconciliation Report to PAO by the HO"
GROUPS = (("Finacle", "D", "E"),
          ("Monthly Cash Account*", "F", "G"),
          ("Difference\n(Finacle - Cash Account)", "H", "I"))
HEAD_ROW, SUB_ROW, FIRST_DATA_ROW = 7, 8, 9
LAST_COL = 9


def _header(ws, ddo_code, ho_name, division, month):
    ws.merge_cells(f"A1:{get_column_letter(LAST_COL)}1")
    ws["A1"] = TITLE
    ws["A1"].font = style.TITLE
    ws["A1"].alignment = style.CENTER

    ws.merge_cells(f"A2:{get_column_letter(LAST_COL)}2")
    ws["A2"] = "(Due Date: 4th of every month)"
    ws["A2"].font = style.SUBTITLE
    ws["A2"].alignment = style.CENTER

    ws.merge_cells(f"A4:{get_column_letter(LAST_COL)}4")
    ws["A4"] = "Table-1"
    ws["A4"].font = style.BOLD
    ws["A4"].alignment = style.CENTER

    ws["A6"] = "DDO Code:"
    ws["A6"].font = style.BOLD
    ws["B6"] = style.safe_text(ddo_code)
    ws["D6"] = "HO:"
    ws["D6"].font = style.BOLD
    ws["E6"] = style.safe_text(ho_name)
    ws["H6"] = "Month:"
    ws["H6"].font = style.BOLD
    ws["I6"] = month.strftime("%b-%Y")

    for ref, label in (("A", "Sl"), ("B", "Account Code"),
                       ("C", "Account Code Description")):
        ws.merge_cells(f"{ref}{HEAD_ROW}:{ref}{SUB_ROW}")
        ws[f"{ref}{HEAD_ROW}"] = label
        style.apply_header(ws[f"{ref}{HEAD_ROW}"])
        style.apply_header(ws[f"{ref}{SUB_ROW}"])

    for label, left, right in GROUPS:
        ws.merge_cells(f"{left}{HEAD_ROW}:{right}{HEAD_ROW}")
        ws[f"{left}{HEAD_ROW}"] = label
        style.apply_header(ws[f"{left}{HEAD_ROW}"])
        style.apply_header(ws[f"{right}{HEAD_ROW}"])
        for ref, side in ((left, "Receipts\n(In Rs.)"),
                          (right, "Payments\n(In Rs.)")):
            ws[f"{ref}{SUB_ROW}"] = side
            style.apply_header(ws[f"{ref}{SUB_ROW}"])

    ws.row_dimensions[HEAD_ROW].height = 40
    ws.row_dimensions[SUB_ROW].height = 30


def _signatures(ws, row, division):
    """The joint signature block the order prescribes."""
    ws.cell(row=row, column=2, value="SBCO In-charge").font = style.BOLD
    ws.cell(row=row, column=7, value="Postmaster").font = style.BOLD
    ws.cell(row=row + 1, column=2, value="HO").font = style.BODY
    ws.cell(row=row + 1, column=7, value="HO").font = style.BODY
    ws.cell(row=row + 3, column=1,
            value="Forwarded to the General Manager(F) / DA(P)").font = style.BODY
    ws.cell(row=row + 4, column=1,
            value=style.safe_text(
                f"Copy to: The S/SPOs, {division} Division for information."
            )).font = style.BODY


def write_annexure_iv_table1(path, result, *, ddo_code: str, ho_name: str,
                             division: str, month: date,
                             include_matched: bool = False) -> Path:
    """Write Table-1. `result` is a Table1Result from annexure.build_table1."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Table-1"
    _header(ws, ddo_code, ho_name, division, month)

    rows = result.rows if include_matched else result.differences
    rows = sorted(rows, key=lambda r: r.account_code)

    for n, item in enumerate(rows, start=1):
        r = FIRST_DATA_ROW + n - 1
        ws.cell(row=r, column=1, value=n).alignment = style.CENTER
        style.write_text(ws, r, 2, item.account_code)
        style.write_text(ws, r, 3, item.description)

        for col, value in ((4, item.finacle_receipt), (5, item.finacle_payment),
                           (6, item.cashbook_receipt), (7, item.cashbook_payment)):
            style.write_money(ws, r, col, value).font = style.BODY

        # difference stays a live formula so the PAO can see the derivation
        for col, finacle, cash, amount in (
                (8, "D", "F", item.difference_receipt),
                (9, "E", "G", item.difference_payment)):
            cell = ws.cell(row=r, column=col, value=f"={finacle}{r}-{cash}{r}")
            cell.font = style.NEGATIVE if amount < ZERO else style.BODY
            cell.number_format = style.CURRENCY
            cell.alignment = style.RIGHT

        for col in range(1, LAST_COL + 1):
            ws.cell(row=r, column=col).border = style.BOX

    last = FIRST_DATA_ROW + len(rows) - 1
    total_row = FIRST_DATA_ROW + len(rows)
    ws.cell(row=total_row, column=3, value="TOTAL").font = style.BOLD
    for col in range(4, LAST_COL + 1):
        letter = get_column_letter(col)
        cell = ws.cell(row=total_row, column=col,
                       value=(f"=SUM({letter}{FIRST_DATA_ROW}:{letter}{last})"
                              if rows else 0))
        cell.font = style.BOLD
        cell.number_format = style.CURRENCY
        cell.alignment = style.RIGHT
        cell.fill = style.TOTAL_FILL
        cell.border = style.BOX
    for col in (1, 2, 3):
        c = ws.cell(row=total_row, column=col)
        c.fill = style.TOTAL_FILL
        c.border = style.BOX

    note = total_row + 2
    ws.cell(row=note, column=1, value=CASH_ACCOUNT_NOTE).font = style.SUBTITLE

    for i, warning in enumerate(result.warnings, start=1):
        ws.cell(row=note + i, column=1, value=warning).font = style.SUBTITLE

    _signatures(ws, note + len(result.warnings) + 3, division)

    style.autosize(ws, [6, 14, 44, 16, 16, 16, 16, 16, 16])
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = f"{HEAD_ROW}:{SUB_ROW}"
    wb.save(Path(path))
    return Path(path)
